import os
import csv
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def create_header_style():
    font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")  # Navy Blue
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Side(border_style="thin", color="CCCCCC")
    border = Border(top=thin_border, left=thin_border, right=thin_border, bottom=thin_border)
    return font, fill, alignment, border


def create_meta_header_style():
    font = Font(name="Calibri", size=10, bold=True, color="1F497D")
    fill = PatternFill(start_color="E9EEF4", end_color="E9EEF4", fill_type="solid")
    alignment = Alignment(horizontal="left", vertical="center")
    thin_border = Side(border_style="thin", color="D0D7DE")
    border = Border(top=thin_border, left=thin_border, right=thin_border, bottom=thin_border)
    return font, fill, alignment, border


def create_data_style(is_alt=False):
    font = Font(name="Calibri", size=10, color="000000")
    fill_color = "F9FAFC" if is_alt else "FFFFFF"
    fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    alignment = Alignment(vertical="top", wrap_text=True)
    thin_border = Side(border_style="thin", color="E0E0E0")
    border = Border(top=thin_border, left=thin_border, right=thin_border, bottom=thin_border)
    return font, fill, alignment, border


def auto_fit_columns(ws, min_width=12, max_width=60):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            # Skip title and metadata rows for width calculation
            if cell.row < 5:
                continue
            val_str = str(cell.value or "")
            if "\n" in val_str:
                lines = val_str.split("\n")
                val_str = max(lines, key=len)
            max_len = max(max_len, len(val_str))
        calc_width = max(min_width, min(max_len + 3, max_width))
        ws.column_dimensions[col_letter].width = calc_width


def generate_tender_workbook(tender_record: Dict[str, Any], output_path: str):
    """Generates an individual formatted Excel workbook for a single tender."""
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tender_Items"
    ws.views.sheetView[0].showGridLines = True

    meta = tender_record.get("metadata", {})
    tno = meta.get("procurement_no") or tender_record.get("file_name", "").replace(".pdf", "")
    order_no = meta.get("order_list_no") or "N/A"
    closing = meta.get("closing_date") or "N/A"
    fee = meta.get("tender_fee") or "N/A"
    bid_val = meta.get("bid_validity") or "N/A"
    bond_val = meta.get("bid_bond_validity") or "N/A"

    m_font, m_fill, m_align, m_border = create_meta_header_style()
    v_font = Font(name="Calibri", size=10, bold=False)

    ws.merge_cells("A1:G1")
    t_cell = ws["A1"]
    t_cell.value = f"STATE PHARMACEUTICALS CORPORATION — SCHEDULE OF REQUIREMENTS"
    t_cell.font = Font(name="Calibri", size=13, bold=True, color="1F497D")
    t_cell.alignment = Alignment(horizontal="left", vertical="center")

    meta_rows = [
        ("Tender / Procurement No:", tno, "Tender Fee:", fee),
        ("MSD Order List No:", order_no, "Closing Date & Time:", closing),
        ("Bid Validity:", bid_val, "Bid Bond Validity:", bond_val)
    ]

    for idx, (lbl1, val1, lbl2, val2) in enumerate(meta_rows, start=2):
        ws.cell(idx, 1, lbl1).font = m_font
        ws.cell(idx, 1).fill = m_fill
        ws.cell(idx, 1).border = m_border

        ws.cell(idx, 2, val1).font = v_font
        ws.cell(idx, 2).border = m_border

        ws.cell(idx, 3, lbl2).font = m_font
        ws.cell(idx, 3).fill = m_fill
        ws.cell(idx, 3).border = m_border

        ws.cell(idx, 4, val2).font = v_font
        ws.cell(idx, 4).border = m_border

    headers = [
        "Item No", "SR Number", "Item Description & Specifications",
        "Quantity", "Delivery Schedule", "Bid Bond Value", "Page"
    ]
    h_font, h_fill, h_align, h_border = create_header_style()

    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(5, col_idx, h_text)
        cell.font = h_font
        cell.fill = h_fill
        cell.alignment = h_align
        cell.border = h_border
    ws.row_dimensions[5].height = 28

    items = tender_record.get("items", [])
    for row_idx, it in enumerate(items, start=6):
        is_alt = (row_idx % 2 == 1)
        d_font, d_fill, d_align, d_border = create_data_style(is_alt)

        row_vals = [
            it.get("item_no", ""),
            it.get("sr_no", ""),
            it.get("description", ""),
            it.get("quantity", ""),
            it.get("delivery_schedule", ""),
            it.get("bid_bond", ""),
            it.get("source_page", "")
        ]

        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row_idx, col_idx, val)
            cell.font = d_font
            cell.fill = d_fill
            cell.border = d_border
            if col_idx in [1, 2, 7]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif col_idx in [4, 5, 6]:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    auto_fit_columns(ws)
    ws.column_dimensions['C'].width = 45
    ws.column_dimensions['D'].width = 16
    ws.column_dimensions['E'].width = 24
    ws.column_dimensions['F'].width = 20

    wb.save(output_path)


def build_consolidated_workbook(tender_records: List[Dict[str, Any]], output_path: str):
    """Builds the comprehensive multi-tab consolidated Excel workbook."""
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default blank sheet

    h_font, h_fill, h_align, h_border = create_header_style()

    # Sheets setup
    ws_all = wb.create_sheet("All_Items")
    ws_sa = wb.create_sheet("DHS_SA_Items")
    ws_ss = wb.create_sheet("DHS_SS_Items")
    ws_cat = wb.create_sheet("Tenders_Catalog")

    for ws in [ws_all, ws_sa, ws_ss, ws_cat]:
        ws.views.sheetView[0].showGridLines = True

    # Item Headers
    item_headers = [
        "Tender No", "Category", "MSD Order List No", "Closing Date",
        "Item No", "SR Number", "Item Description & Specifications",
        "Quantity", "Delivery Schedule", "Bid Bond Value", "Page", "Source File"
    ]
    for ws in [ws_all, ws_sa, ws_ss]:
        for c_idx, h in enumerate(item_headers, start=1):
            cell = ws.cell(1, c_idx, h)
            cell.font = h_font
            cell.fill = h_fill
            cell.alignment = h_align
            cell.border = h_border
        ws.row_dimensions[1].height = 28

    # Catalog Headers
    cat_headers = [
        "Tender No", "Category", "MSD Order List No", "Closing Date",
        "Item Count", "Tender Fee", "Bid Validity", "Bid Bond Validity",
        "Page Count", "Status", "Source File"
    ]
    for c_idx, h in enumerate(cat_headers, start=1):
        cell = ws_cat.cell(1, c_idx, h)
        cell.font = h_font
        cell.fill = h_fill
        cell.alignment = h_align
        cell.border = h_border
    ws_cat.row_dimensions[1].height = 28

    row_counters = {ws_all: 2, ws_sa: 2, ws_ss: 2, ws_cat: 2}

    for t_idx, tr in enumerate(tender_records):
        meta = tr.get("metadata", {})
        fname = tr.get("file_name", "")
        tno = meta.get("procurement_no") or fname.replace(".pdf", "").replace("_", "/")
        category = "DHS/SS" if "DHS/SS" in tno.upper() or "DHS_SS" in fname else "DHS/SA"
        order_no = meta.get("order_list_no") or ""
        closing = meta.get("closing_date") or ""
        items = tr.get("items", [])

        # Catalog Entry
        cat_ws = ws_cat
        c_row = row_counters[cat_ws]
        is_alt = (c_row % 2 == 1)
        d_font, d_fill, _, d_border = create_data_style(is_alt)

        cat_vals = [
            tno, category, order_no, closing, len(items),
            meta.get("tender_fee", ""), meta.get("bid_validity", ""),
            meta.get("bid_bond_validity", ""), meta.get("page_count", 0),
            tr.get("status", "SUCCESS"), fname
        ]
        for col_idx, val in enumerate(cat_vals, start=1):
            cell = cat_ws.cell(c_row, col_idx, val)
            cell.font = d_font
            cell.fill = d_fill
            cell.border = d_border
            if col_idx in [1, 2, 4, 5, 9, 10]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
        row_counters[cat_ws] += 1

        # Items Entries
        target_sub_ws = ws_ss if category == "DHS/SS" else ws_sa
        for it in items:
            for dest_ws in [ws_all, target_sub_ws]:
                r = row_counters[dest_ws]
                is_a = (r % 2 == 1)
                df, dfl, _, db = create_data_style(is_a)

                row_vals = [
                    tno, category, order_no, closing,
                    it.get("item_no", ""), it.get("sr_no", ""),
                    it.get("description", ""), it.get("quantity", ""),
                    it.get("delivery_schedule", ""), it.get("bid_bond", ""),
                    it.get("source_page", ""), fname
                ]
                for col_idx, val in enumerate(row_vals, start=1):
                    cell = dest_ws.cell(r, col_idx, val)
                    cell.font = df
                    cell.fill = dfl
                    cell.border = db
                    if col_idx in [1, 2, 4, 5, 6, 11]:
                        cell.alignment = Alignment(horizontal="center", vertical="top")
                    elif col_idx in [8, 9, 10]:
                        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                row_counters[dest_ws] += 1

    for ws in [ws_all, ws_sa, ws_ss]:
        auto_fit_columns(ws)
        ws.column_dimensions['G'].width = 45
        ws.column_dimensions['H'].width = 16
        ws.column_dimensions['I'].width = 24
        ws.column_dimensions['J'].width = 20

    auto_fit_columns(ws_cat)
    wb.save(output_path)


def generate_csv(tender_records: List[Dict[str, Any]], output_path: str):
    """Generates a clean consolidated CSV file of all extracted line items."""
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    headers = [
        "tender_no", "category", "order_list_no", "closing_date",
        "item_no", "sr_no", "description", "quantity",
        "delivery_schedule", "bid_bond", "page", "source_file"
    ]

    rows = []
    for tr in tender_records:
        meta = tr.get("metadata", {})
        fname = tr.get("file_name", "")
        tno = meta.get("procurement_no") or fname.replace(".pdf", "").replace("_", "/")
        category = "DHS/SS" if "DHS/SS" in tno.upper() or "DHS_SS" in fname else "DHS/SA"
        order_no = meta.get("order_list_no") or ""
        closing = meta.get("closing_date") or ""

        for it in tr.get("items", []):
            rows.append({
                "tender_no": tno,
                "category": category,
                "order_list_no": order_no,
                "closing_date": closing,
                "item_no": it.get("item_no", ""),
                "sr_no": it.get("sr_no", ""),
                "description": it.get("description", ""),
                "quantity": it.get("quantity", ""),
                "delivery_schedule": it.get("delivery_schedule", ""),
                "bid_bond": it.get("bid_bond", ""),
                "page": it.get("source_page", ""),
                "source_file": fname
            })

    with open(output_path, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)
